"""Renders tailored resume and cover letter to plain text, DOCX, and PDF.

Both documents are built as a list of blocks, then rendered by each backend, so the three formats
never drift apart.
"""
from __future__ import annotations

import io
from datetime import date
from xml.sax.saxutils import escape

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import ListFlowable, ListItem, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from .models import Job
from .profile_parser import Profile
from .tailoring import TailoredDocs

Block = tuple  # (kind, *args): name | contact | heading | para | role(left, right) | bullet | spacer


def _contact_line(p: Profile) -> str:
    c = p.contact
    return " | ".join(filter(None, [c.get("location"), c.get("email"), c.get("phone"), c.get("linkedin"), c.get("github")]))


def resume_blocks(docs: TailoredDocs, p: Profile) -> list[Block]:
    certs = {c["name"]: c for c in p.certifications}
    cert_blocks = [("heading", "Certifications")] + [
        ("bullet", " — ".join(filter(None, [name, certs[name]["issuer"], certs[name]["date"]])))
        for name in docs.certification_order
    ]
    blocks: list[Block] = [
        ("name", p.contact.get("name", "")),
        ("contact", _contact_line(p)),
        ("heading", "Professional Summary"),
        ("para", docs.summary),
    ]
    if docs.certifications_near_top:
        blocks += cert_blocks
    blocks.append(("heading", "Skills"))
    blocks += [("para", f"**{g['category']}:** {', '.join(g['items'])}") for g in docs.skills]

    blocks.append(("heading", "Experience"))
    roles = {e.key: e for e in p.experiences}
    for item in docs.experience:
        r = roles[item["key"]]
        blocks.append(("role", f"{r.title} — {r.company}, {r.location}", r.dates))
        blocks += [("bullet", b) for b in item["bullets"]]

    if not docs.certifications_near_top:
        blocks += cert_blocks
    blocks.append(("heading", "Education"))
    blocks += [("para", e) for e in p.education]
    return blocks


def cover_letter_blocks(docs: TailoredDocs, p: Profile, job: Job) -> list[Block]:
    blocks: list[Block] = [
        ("name", p.contact.get("name", "")),
        ("contact", _contact_line(p)),
        ("spacer",),
        ("para", date.today().strftime("%B %d, %Y").replace(" 0", " ")),
        ("para", f"Re: {job.title}"),
        ("para", f"Dear Hiring Team at {job.company}," if job.company else "Dear Hiring Team,"),
    ]
    blocks += [("para", para.strip()) for para in docs.cover_letter_body.split("\n\n") if para.strip()]
    blocks += [("para", "Sincerely,"), ("para", p.contact.get("name", ""))]
    return blocks


# ---------------------------------------------------------------------------- text
def to_text(blocks: list[Block]) -> str:
    lines = []
    for kind, *args in blocks:
        if kind == "heading":
            lines += ["", args[0].upper()]
        elif kind == "role":
            lines += ["", f"{args[0]} | {args[1]}"]
        elif kind == "bullet":
            lines.append(f"- {args[0]}")
        elif kind == "spacer":
            lines.append("")
        else:
            lines.append(args[0].replace("**", ""))
    return "\n".join(lines).strip() + "\n"


# ---------------------------------------------------------------------------- docx
def _docx_runs(par, text: str, size: float = 10.5):
    # supports **bold** segments only
    for i, part in enumerate(text.split("**")):
        run = par.add_run(part)
        run.bold = i % 2 == 1
        run.font.size = Pt(size)


def to_docx(blocks: list[Block]) -> bytes:
    d = Document()
    for s in d.sections:
        s.left_margin = s.right_margin = s.top_margin = s.bottom_margin = Pt(50)
    d.styles["Normal"].font.name = "Calibri"
    for kind, *args in blocks:
        if kind == "name":
            par = d.add_paragraph()
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = par.add_run(args[0])
            run.bold, run.font.size = True, Pt(18)
        elif kind == "contact":
            par = d.add_paragraph()
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            _docx_runs(par, args[0], 9.5)
        elif kind == "heading":
            par = d.add_paragraph()
            par.paragraph_format.space_before = Pt(10)
            run = par.add_run(args[0].upper())
            run.bold, run.font.size = True, Pt(11.5)
        elif kind == "role":
            par = d.add_paragraph()
            par.paragraph_format.space_before = Pt(6)
            _docx_runs(par, f"**{args[0]}**  |  {args[1]}")
        elif kind == "bullet":
            _docx_runs(d.add_paragraph(style="List Bullet"), args[0])
        elif kind == "spacer":
            d.add_paragraph()
        else:
            _docx_runs(d.add_paragraph(), args[0])
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------- pdf
_S = {
    "name": ParagraphStyle("name", fontName="Helvetica-Bold", fontSize=17, leading=21, alignment=TA_CENTER),
    "contact": ParagraphStyle("contact", fontName="Helvetica", fontSize=8.5, leading=11, alignment=TA_CENTER),
    "heading": ParagraphStyle("heading", fontName="Helvetica-Bold", fontSize=10.5, leading=13, spaceBefore=9, spaceAfter=3),
    "para": ParagraphStyle("para", fontName="Helvetica", fontSize=9.5, leading=12.5, spaceAfter=4),
    "role": ParagraphStyle("role", fontName="Helvetica-Bold", fontSize=9.5, leading=12),
    "dates": ParagraphStyle("dates", fontName="Helvetica", fontSize=9, leading=12, alignment=2),
}


def _rl(text: str) -> str:
    parts = escape(text).split("**")
    return "".join(f"<b>{p}</b>" if i % 2 else p for i, p in enumerate(parts))


def to_pdf(blocks: list[Block]) -> bytes:
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=17 * mm, rightMargin=17 * mm, topMargin=15 * mm, bottomMargin=15 * mm)
    flow, bullets = [], []

    def flush():
        if bullets:
            flow.append(ListFlowable([ListItem(Paragraph(_rl(b), _S["para"]), leftIndent=10) for b in bullets],
                                     bulletType="bullet", start="•", leftIndent=10, bulletFontSize=8))
            bullets.clear()

    for kind, *args in blocks:
        if kind == "bullet":
            bullets.append(args[0])
            continue
        flush()
        if kind == "role":
            t = Table([[Paragraph(_rl(args[0]), _S["role"]), Paragraph(_rl(args[1]), _S["dates"])]],
                      colWidths=[doc.width * 0.68, doc.width * 0.32])
            t.setStyle(TableStyle([("LEFTPADDING", (0, 0), (-1, -1), 0), ("RIGHTPADDING", (0, 0), (-1, -1), 0),
                                   ("TOPPADDING", (0, 0), (-1, -1), 5), ("VALIGN", (0, 0), (-1, -1), "TOP")]))
            flow.append(t)
        elif kind == "spacer":
            flow.append(Spacer(1, 8))
        elif kind == "heading":
            flow.append(Paragraph(_rl(args[0].upper()), _S["heading"]))
        else:
            flow.append(Paragraph(_rl(args[0]), _S[kind if kind in _S else "para"]))
    flush()
    doc.build(flow)
    return buf.getvalue()


def render_all(docs: TailoredDocs, profile: Profile, job: Job) -> dict:
    r, c = resume_blocks(docs, profile), cover_letter_blocks(docs, profile, job)
    return {
        "resume_text": to_text(r), "resume_pdf": to_pdf(r), "resume_docx": to_docx(r),
        "cover_text": to_text(c), "cover_pdf": to_pdf(c), "cover_docx": to_docx(c),
    }
